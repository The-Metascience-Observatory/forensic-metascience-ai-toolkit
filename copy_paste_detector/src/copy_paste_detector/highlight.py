"""
Highlighted-workbook artifact: mark detect_raw_duplication's match groups in a
copy of the source spreadsheet so a human can verify them in one glance.

Best-effort by contract, exactly like `report_render`: a missing dependency, a
vanished source file, or an unwritable path returns a structured `error` and
NEVER raises — the artifact is a convenience on top of the findings, not a
step the pipeline may fail on.

Rendering rules:
  - one light fill color per match group, cycled from a fixed palette, so the
    two (or N) occurrences of a group can be paired by eye;
  - a thin black border boxes every marked cell;
  - the FIRST cell of each group carries a comment naming the group kind and
    every other location, so the partner rows can be found without scrolling;
  - overlap avoidance: a group whose cells are already styled by an earlier
    (more severe) group is skipped and counted, never painted over — a muddled
    two-color cell is worse than a skipped group.

Only .xlsx output. A .csv/.tsv source is rehydrated into a fresh workbook (one
sheet) so it can carry styling; .docx has no cell coordinates and is refused.
"""

from __future__ import annotations

import csv
from pathlib import Path
from typing import Any

#: Light ARGB fills, distinct enough to pair occurrences by eye. Severity is
#: NOT encoded in color (the report carries severity); color = group identity.
_PALETTE = (
    "FFFFF2AC",
    "FFD3F0D1",
    "FFD6E4FF",
    "FFFFD9D1",
    "FFE8D9FF",
    "FFD1F0EC",
    "FFFFE1F0",
    "FFEDEFC8",
    "FFDCEBDD",
    "FFF0E0C8",
    "FFDDE7F5",
    "FFF5DDDD",
)

_MAX_GROUPS = 50  # highlighting hundreds of groups helps nobody


def _group_cells(group: dict) -> list[dict]:
    """Every cell reference a duplicate group carries, whatever its kind."""
    if group.get("kind") == "column_sequence":
        return list(group.get("cells_a", ())) + list(group.get("cells_b", ()))
    return list(group.get("cells", ()))


def _describe(group: dict) -> str:
    kind = group.get("kind", "match")
    sev = group.get("severity", "")
    locs = ", ".join(
        f"{c.get('sheet', '?')}!{c.get('cell', '?')}" for c in _group_cells(group)[:16]
    )
    return f"detect_raw_duplication: {kind} ({sev}); matching cells: {locs}"


def highlight_workbook(
    src_path: str | Path,
    duplicate_groups: list[dict],
    out_path: str | Path,
) -> dict[str, Any]:
    """Write a highlighted .xlsx copy of `src_path` marking each match group.

    Returns {"path", "n_groups_highlighted", "n_groups_skipped"} on success or
    {"error": ...} on any failure. Never raises.
    """
    try:
        return _highlight(Path(src_path), duplicate_groups, Path(out_path))
    except Exception as exc:  # best-effort contract: never break the caller
        return {"error": f"could not write highlighted workbook: {type(exc).__name__}: {exc}"}


def _highlight(src: Path, groups: list[dict], out: Path) -> dict[str, Any]:
    if not src.exists():
        return {"error": f"source file not found: {src}"}
    groups = [g for g in groups if _group_cells(g)]
    if not groups:
        return {"error": "no duplicate groups with cell references to highlight"}

    try:
        import openpyxl
        from openpyxl.comments import Comment
        from openpyxl.styles import Border, PatternFill, Side
    except ImportError:
        return {"error": "openpyxl is not installed; cannot write .xlsx"}

    suffix = src.suffix.lower()
    if suffix in (".xlsx", ".xlsm"):
        # data_only=False keeps formulas; styling does not touch values.
        book = openpyxl.load_workbook(str(src))
        sheet_for = {ws.title: ws for ws in book.worksheets}
        fallback = None
    elif suffix in (".csv", ".tsv"):
        book = openpyxl.Workbook()
        ws = book.active
        ws.title = src.stem[:31] or "data"
        delimiter = "\t" if suffix == ".tsv" else ","
        with src.open(encoding="utf-8", errors="replace", newline="") as handle:
            for row in csv.reader(handle, delimiter=delimiter):
                ws.append(row)
        sheet_for = {ws.title: ws}
        fallback = ws  # single sheet: every ref lands here whatever its label
    else:
        return {
            "error": f"cannot highlight {suffix} (no cell coordinates); "
            "only .xlsx/.xlsm/.csv/.tsv sources"
        }

    thin = Side(style="thin", color="FF000000")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)

    styled: set[tuple[str, str]] = set()
    highlighted = 0
    skipped = 0
    for g_idx, group in enumerate(groups[:_MAX_GROUPS]):
        cells = _group_cells(group)
        targets = []
        for ref in cells:
            ws = sheet_for.get(ref.get("sheet", ""), fallback)
            if ws is None:
                targets = []
                break
            targets.append((ws, ref.get("cell", "")))
        if not targets:
            skipped += 1
            continue
        if any((ws.title, addr) in styled for ws, addr in targets):
            skipped += 1
            continue
        fill = PatternFill("solid", fgColor=_PALETTE[g_idx % len(_PALETTE)])
        for ws, addr in targets:
            cell = ws[addr]
            cell.fill = fill
            cell.border = box
            styled.add((ws.title, addr))
        first_ws, first_addr = targets[0]
        first_ws[first_addr].comment = Comment(_describe(group), "forensic-metascience-ai-toolkit")
        highlighted += 1
    if len(groups) > _MAX_GROUPS:
        skipped += len(groups) - _MAX_GROUPS

    out.parent.mkdir(parents=True, exist_ok=True)
    book.save(str(out))
    return {
        "path": str(out),
        "n_groups_highlighted": highlighted,
        "n_groups_skipped": skipped,
    }

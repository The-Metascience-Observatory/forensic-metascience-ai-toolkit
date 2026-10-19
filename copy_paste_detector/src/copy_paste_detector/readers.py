"""Load raw tables into numeric segments and retain cell references."""

from __future__ import annotations

import csv
import math
from collections import defaultdict
from dataclasses import asdict, dataclass
from pathlib import Path
from typing import Any, Optional

from .constants import (
    _HAS_ALPHA,
    _MAX_COLS,
    _MAX_ROWS,
    _MAX_SHEETS,
    _THOUSANDS,
)
from .scoring import (
    _canon,
)
from .severity import (
    INDETERMINATE,
)


@dataclass(frozen=True)
class DetectionLimits:
    """Zero means unlimited. Limits describe input, not scientific applicability."""

    max_rows: int = _MAX_ROWS
    max_columns: int = _MAX_COLS
    max_sheets: int = _MAX_SHEETS

    def __post_init__(self):
        if any(
            isinstance(v, bool) or not isinstance(v, int) or v < 0 for v in asdict(self).values()
        ):
            raise ValueError("Reader limits must be nonnegative integers (0 = unlimited)")


class _Segment:
    """One logical table: a sheet, or one block of a sheet that stacks
    several tables separated by blank rows + a fresh header row."""

    __slots__ = (
        "sheet",
        "label",
        "col_names",
        "rows",
        "n_numeric",
        "n_text",
        "source",
        "raw",
        "exclude_columns",
    )

    def __init__(self, sheet: str, label: str):
        self.source = None
        self.raw = {}
        self.exclude_columns = []
        self.sheet = sheet
        self.label = label
        self.col_names: list[str] = []
        #: list of (original 1-based row number, {col_index: canonical float})
        self.rows: list[tuple[int, dict[int, float]]] = []
        self.n_numeric = 0
        #: {col_index: count of non-empty cells that did NOT parse as numbers}.
        #: A column of "80.0.3987.163" browser builds is text, but the few
        #: entries shaped like "75.0" parse -- leaving a numeric minority that
        #: reads as high-precision data. Counting the rejects is what lets
        #: _screen_columns recognise the column as text and drop it whole.
        self.n_text: dict[int, int] = defaultdict(int)


def _as_number(value: Any) -> Optional[float]:
    """A cell's numeric reading, or None. Dates/datetimes, strings with
    letters, and empties are never numbers."""
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        f = float(value)
        return f if math.isfinite(f) else None  # inf/NaN would crash _canon
    if hasattr(value, "year") and hasattr(value, "month"):  # date/datetime
        return None
    if isinstance(value, str):
        s = value.strip()
        if "," in s:
            if not _THOUSANDS.match(s):
                return None  # a list ("1,2,4,3"), not a number
            s = s.replace(",", "")
        if not s or _HAS_ALPHA.search(s):
            return None
        try:
            f = float(s)
        except ValueError:
            return None
        # A 400-digit cell parses to float('inf'); it is not a number here.
        return f if math.isfinite(f) else None
    return None


def _cell_label(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _segment_rows(
    sheet_name: str,
    raw_rows: list[tuple[int, list[Any]]],
    caveats: list[str],
) -> list[_Segment]:
    """Split one sheet's rows into logical segments.

    A supplement sheet sometimes stacks several tables: data, blank row(s),
    a fresh header row, more data. Mixing the blocks would compare values from
    DIFFERENT variables under one column identity, so a new segment starts when
    a blank gap is followed by a text-only row with 2+ labels (a one-cell text
    row is a section label and stays inside its segment).
    """
    segments: list[_Segment] = []
    current: Optional[_Segment] = None
    pending_header: Optional[list[str]] = None
    saw_blank_gap = False
    seg_count = 0

    def _new_segment(header: Optional[list[str]]) -> _Segment:
        nonlocal seg_count
        seg_count += 1
        label = sheet_name if seg_count == 1 else f"{sheet_name} (block {seg_count})"
        seg = _Segment(sheet_name, label)
        seg.col_names = header or []
        return seg

    for row_no, cells in raw_rows:
        numbers = {
            i: _canon(v)
            for i, v in ((i, _as_number(c)) for i, c in enumerate(cells))
            if v is not None
        }
        labels = [_cell_label(c) for c in cells]
        non_empty = sum(1 for s in labels if s)

        if non_empty == 0:
            saw_blank_gap = current is not None and bool(current.rows)
            continue

        if not numbers:
            # Text-only row: header material. Starts a NEW segment only after
            # a blank gap below existing data; otherwise it extends the
            # current header or is a section label inside the segment.
            if saw_blank_gap and non_empty >= 2:
                if current is not None and current.rows:
                    segments.append(current)
                current = None
                pending_header = labels
            elif current is None:
                pending_header = (
                    labels
                    if pending_header is None
                    else [
                        " ".join(filter(None, pair)).strip()
                        for pair in __import__("itertools").zip_longest(
                            pending_header, labels, fillvalue=""
                        )
                    ]
                )
            saw_blank_gap = False
            continue

        saw_blank_gap = False
        if current is None:
            current = _new_segment(pending_header)
            pending_header = None
        current.rows.append((row_no, numbers))
        current.raw[row_no] = labels
        current.n_numeric += len(numbers)
        for i, s in enumerate(labels):
            if s and i not in numbers:
                current.n_text[i] += 1

    if current is not None and current.rows:
        segments.append(current)
    if seg_count > 1 and len(segments) > 1:
        caveats.append(
            f"sheet '{sheet_name}' stacks {len(segments)} tables (blank gap + "
            "fresh header); each block was analyzed as its own table"
        )
    return segments


def _load_segments(
    path: Path, sheet: Optional[str], caveats: list[str], *, limits=None, coverage=None
) -> list[_Segment] | dict:
    """All logical segments of the file, or a structured error dict."""
    limits = limits or DetectionLimits()
    coverage = coverage if coverage is not None else {}
    coverage.update(limits=asdict(limits), omitted=[], missing_formula_cache=[], status="complete")

    def omit(reason):
        coverage["omitted"].append(reason)
        coverage["status"] = "partial"
        caveats.append(reason)

    def bounded(rows, label):
        raw = []
        column_warning = False
        for i, row in enumerate(rows, 1):
            if limits.max_rows and i > limits.max_rows:
                omit(
                    f"{label}: only the first {limits.max_rows} rows analyzed; additional rows exist"
                )
                break
            if limits.max_columns and len(row) > limits.max_columns:
                if not column_warning:
                    omit(
                        f"{label}: only the first {limits.max_columns} columns analyzed; additional columns exist"
                    )
                    column_warning = True
                row = row[: limits.max_columns]
            raw.append((i, list(row)))
        return raw

    suffix = path.suffix.lower()

    if suffix in (".csv", ".tsv"):
        delimiter = "\t" if suffix == ".tsv" else ","
        with path.open(encoding="utf-8-sig", errors="strict", newline="") as handle:
            raw = bounded(csv.reader(handle, delimiter=delimiter), path.name)
        return _segment_rows(path.name, raw, caveats)

    if suffix in (".xlsx", ".xlsm"):
        try:
            import openpyxl
        except ImportError:
            return {
                "error": "openpyxl is not installed; cannot read .xlsx",
                "severity": INDETERMINATE,
            }
        book = openpyxl.load_workbook(str(path), data_only=True, read_only=True)
        try:
            names = [ws.title for ws in book.worksheets]
            if sheet is not None and sheet not in names:
                return {
                    "error": f"sheet {sheet!r} not found in {path.name}; sheets are {names}",
                    "severity": INDETERMINATE,
                }
            selected = [ws for ws in book.worksheets if sheet is None or ws.title == sheet]
            if limits.max_sheets and len(selected) > limits.max_sheets:
                dropped = [ws.title for ws in selected[limits.max_sheets :]]
                omit(
                    f"only the first {limits.max_sheets} sheets were analyzed; NOT analyzed: {dropped}"
                )
                selected = selected[: limits.max_sheets]
            segments: list[_Segment] = []
            for ws in selected:
                raw = bounded(ws.iter_rows(values_only=True), f"sheet '{ws.title}'")
                segments.extend(_segment_rows(ws.title, raw, caveats))
            # A second streaming pass exposes formulas with no cached value.
            # These are not blanks or negative measurements. Never calculate them.
            formulas = openpyxl.load_workbook(str(path), data_only=False, read_only=True)
            try:
                selected_names = {ws.title for ws in selected}
                for ws in formulas.worksheets:
                    if ws.title not in selected_names:
                        continue
                    cached = book[ws.title].iter_rows(values_only=True)
                    for row_no, (formula_row, cached_row) in enumerate(
                        zip(ws.iter_rows(), cached), 1
                    ):
                        if limits.max_rows and row_no > limits.max_rows:
                            break
                        for col, (cell, value) in enumerate(zip(formula_row, cached_row)):
                            if limits.max_columns and col >= limits.max_columns:
                                break
                            if cell.data_type == "f" and value is None:
                                coverage["missing_formula_cache"].append(
                                    {"sheet": ws.title, "cell": cell.coordinate}
                                )
                if coverage["missing_formula_cache"]:
                    omit(
                        f"{len(coverage['missing_formula_cache'])} formula cells have no cached value; not evaluated"
                    )
            finally:
                formulas.close()
            return segments
        finally:
            book.close()

    if suffix == ".docx":
        # No cell coordinates in a Word table anyway; reuse the shared reader.
        #
        # THE RAW EXTRACTOR, NOT `extract_supplement`. The latter caps its
        # payload at 120k characters and hands back an INDEX instead -- five
        # rows and twelve columns per table, so a reader can CHOOSE which sheet
        # to ask for. That is right for a model with a context window and wrong
        # here: this tool reads the grid programmatically, and analysing a
        # 5-row preview of a 17-row table would report a clean result on data
        # it never saw. A 263k-character supplement in the corpus did exactly
        # that (KeyError on the missing 'data' key -- a loud failure, and
        # lucky: silently scanning the preview is the version that accuses).
        from .docx_reader import extract_docx_tables

        result = extract_docx_tables(path)
        if result.get("error"):
            return {"error": result["error"], "severity": INDETERMINATE}
        segments = []
        for t in result["tables"]:
            grid = t.get("rows") or t.get("data") or []
            raw = [(i, row) for i, row in enumerate(grid, start=1)]
            segments.extend(_segment_rows(t["label"], raw, caveats))
        return segments

    from .formats import UNREADABLE

    if suffix in UNREADABLE:
        return {"error": f"cannot read {suffix}: {UNREADABLE[suffix]}", "severity": INDETERMINATE}
    return {
        "error": f"unsupported raw-data format {suffix!r} "
        "(readable: .xlsx, .xlsm, .csv, .tsv, .docx)",
        "severity": INDETERMINATE,
    }


def _col_letter(idx: int) -> str:
    letters = ""
    idx += 1
    while idx:
        idx, rem = divmod(idx - 1, 26)
        letters = chr(65 + rem) + letters
    return letters


def _cell_ref(seg: _Segment, row_no: int, col: int) -> dict:
    ref = {"sheet": seg.sheet, "cell": f"{_col_letter(col)}{row_no}", "row": row_no, "col": col}
    if seg.source:
        ref.update(seg.source)
        ref["segment"] = seg.label
        ref["column"] = seg.col_names[col] if col < len(seg.col_names) else ""
        ref["value"] = (
            _as_number(seg.raw[row_no][col]) if col < len(seg.raw.get(row_no, [])) else None
        )
    return ref

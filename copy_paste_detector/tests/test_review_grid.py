import csv

from openpyxl import Workbook

from copy_paste_detector.review_grid import evidence_grid


def test_wide_late_excel_cells_remain_visible(tmp_path):
    path = tmp_path / "wide.xlsx"
    book = Workbook()
    book.active.cell(1, 102, "mass_mg")
    book.active.cell(1501, 102, 123.456789)
    book.save(path)
    grid = evidence_grid(path, "Sheet", [{"sheet": "Sheet", "row": 1501, "col": 101}])
    mark = grid["highlights"][0]
    assert grid["rows"][mark["r"]][mark["c"]] == "123.456789"
    assert grid["row_numbers"][mark["r"]] == 1501
    assert grid["column_labels"][mark["c"]] == "CX"
    assert len(grid["rows"]) <= 4 and grid["n_cols"] <= 4
    assert "1 of 1" in grid["preview_note"]


def test_large_candidate_preview_announces_omitted_marks(tmp_path):
    path = tmp_path / "data.csv"
    with path.open("w") as f:
        csv.writer(f).writerows([["mass"]] + [[n + 0.23456] for n in range(100)])
    grid = evidence_grid(
        path, highlights=[{"row": n, "col": 0} for n in range(2, 102)], max_rows=10
    )
    assert grid["truncated"]
    assert len(grid["highlights"]) == 9
    assert "9 of 100" in grid["preview_note"]

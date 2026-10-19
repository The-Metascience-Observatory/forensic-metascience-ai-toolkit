"""Bounded source previews around evidence, preserving original row/column numbers."""

from __future__ import annotations

import csv
from pathlib import Path

from .readers import _col_letter


def evidence_grid(path, sheet=None, highlights=(), *, max_rows=500, max_columns=100):
    path = Path(path)
    book = None
    try:
        if path.suffix.lower() in {".xlsx", ".xlsm"}:
            import openpyxl

            book = openpyxl.load_workbook(path, read_only=True, data_only=True)
            names = book.sheetnames
            active = sheet if sheet in names else names[0]
            iterator = book[active].iter_rows(values_only=True)
            refs = [h for h in highlights if not h.get("sheet") or h["sheet"] == active]
        elif path.suffix.lower() in {".csv", ".tsv"}:
            names, active = [path.stem], path.stem
            book = path.open(encoding="utf-8-sig", newline="")
            iterator = csv.reader(book, delimiter="\t" if path.suffix.lower() == ".tsv" else ",")
            refs = list(highlights)
        else:
            raise ValueError("Unsupported evidence grid format")
        coords = sorted(
            {
                (h["row"], h["col"])
                for h in refs
                if isinstance(h.get("row"), int)
                and h["row"] > 0
                and isinstance(h.get("col"), int)
                and h["col"] >= 0
            }
        )
        # Select matched coordinates first; neighboring context uses remaining capacity.
        wanted_rows = sorted({r for r, _ in coords})[: max_rows - 1]
        wanted_cols = sorted({c for _, c in coords})[:max_columns]
        rows_set, cols_set = {1, *wanted_rows}, set(wanted_cols)
        for r in wanted_rows:
            for n in (r - 1, r + 1):
                if n > 0 and len(rows_set) < max_rows:
                    rows_set.add(n)
        for c in [0] + [n for c in wanted_cols for n in (c - 1, c + 1) if n >= 0]:
            if len(cols_set) < max_columns:
                cols_set.add(c)
        if not coords:
            rows_set = set(range(1, min(max_rows, 100) + 1))
            cols_set = set(range(min(max_columns, 20)))
        cols = sorted(cols_set)
        rows, row_numbers = [], []
        for n, row in enumerate(iterator, 1):
            if n > max(rows_set):
                break
            if n in rows_set:
                row_numbers.append(n)
                rows.append(["" if c >= len(row) or row[c] is None else str(row[c]) for c in cols])
        row_map, col_map = (
            {r: i for i, r in enumerate(row_numbers)},
            {c: i for i, c in enumerate(cols)},
        )
        marks = [
            {"r": row_map[r], "c": col_map[c]} for r, c in coords if r in row_map and c in col_map
        ]
        return {
            "rows": rows,
            "row_numbers": row_numbers,
            "column_labels": [_col_letter(c) for c in cols],
            "sheets": names,
            "sheet": active,
            "n_rows": len(rows),
            "n_cols": len(cols),
            "highlights": marks,
            "truncated": len(marks) < len(coords),
            "error": None,
            "preview_note": f"Evidence preview with original coordinates; {len(marks)} of {len(coords)} matched cells shown. Download the source for the entire table.",
        }
    except (OSError, ValueError, KeyError, IndexError) as exc:
        return {"error": f"{type(exc).__name__}: {exc}", "rows": [], "sheets": []}
    finally:
        if book is not None:
            book.close()

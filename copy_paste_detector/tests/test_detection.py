"""detect_raw_duplication: planted positives and one test per FP guard.

The guards matter more than the positives: every false-positive class this
tool suppresses in CODE (shared covariates, ID runs, transforms, years,
zeros, chance-level repeats, pooled data) is pinned here, because a raw-data
duplication detector that accuses innocent structure is worse than none.
"""

from __future__ import annotations

import csv
import math
import random

import pytest

from copy_paste_detector import detect_raw_duplication
from copy_paste_detector.scoring import _info_digits, _sig_digits

openpyxl = pytest.importorskip("openpyxl")


# ---------------------------------------------------------------------------
# fixtures
# ---------------------------------------------------------------------------


def _write_csv(path, rows):
    with path.open("w", newline="") as fh:
        csv.writer(fh).writerows(rows)
    return path


def _clean_rows(n=60, seed=7):
    rng = random.Random(seed)
    rows = [["id", "group", "weight_kg", "cortisol_nmol"]]
    for i in range(1, n + 1):
        rows.append(
            [i, (i % 3) + 1, round(rng.uniform(52, 96), 2), round(rng.uniform(180, 620), 3)]
        )
    return rows


@pytest.fixture()
def clean_csv(tmp_path):
    return _write_csv(tmp_path / "clean.csv", _clean_rows())


# ---------------------------------------------------------------------------
# information content
# ---------------------------------------------------------------------------


class TestInfoDigits:
    def test_significant_digits_strip_zeros(self):
        assert _sig_digits(123.46) == 5
        assert _sig_digits(100.0) == 1
        assert _sig_digits(0.030) == 1
        assert _sig_digits(0.0) == 0

    def test_years_carry_no_information(self):
        assert _info_digits(2015.0) == 1
        assert _info_digits(1997) == 1
        # A year-magnitude value with decimals is NOT a year.
        assert _info_digits(2015.37) > 1

    def test_repeating_fraction_scored_as_source(self):
        # 13/3 prints as 4.333... at full float precision but carries ~2 digits.
        assert _info_digits(13 / 3) <= 2

    def test_square_root_scored_as_radicand(self):
        assert _info_digits(math.sqrt(0.30)) <= 2

    def test_natural_log_scored_as_source(self):
        assert _info_digits(math.log(2.0)) <= 2

    def test_ordinary_precision_keeps_its_digits(self):
        assert _info_digits(413.762) == 6


# ---------------------------------------------------------------------------
# planted positives
# ---------------------------------------------------------------------------


class TestPositives:
    def test_planted_duplicate_row_pair(self, tmp_path):
        rows = _clean_rows()
        rows[47] = [47, rows[12][1], rows[12][2], rows[12][3]]
        res = detect_raw_duplication(str(_write_csv(tmp_path / "d.csv", rows)))
        assert res["severity"] in ("suspicious", "highly_suspicious")
        pairs = [g for g in res["duplicate_groups"] if g["kind"] == "row_pair"]
        assert pairs, "the duplicated row pair was not reported"
        got_rows = {r["row"] for r in pairs[0]["rows"]}
        assert got_rows == {13, 48}  # 1-based, header is row 1
        # occurrence counts travel with the evidence
        assert all("occurrences_in_file" in s for s in pairs[0]["shared"])

    def test_copied_vertical_run(self, tmp_path):
        rows = _clean_rows()
        for k in range(5):
            rows[20 + k][2] = rows[40 + k][3]
        res = detect_raw_duplication(str(_write_csv(tmp_path / "s.csv", rows)))
        seqs = [g for g in res["duplicate_groups"] if g["kind"] == "column_sequence"]
        assert seqs and seqs[0]["length"] >= 5
        assert res["severity"] == "highly_suspicious"

    def test_repeated_high_precision_value(self, tmp_path):
        rows = _clean_rows()
        rows[5][3] = 413.762
        rows[33][2] = 413.762
        res = detect_raw_duplication(str(_write_csv(tmp_path / "v.csv", rows)))
        vals = [g for g in res["duplicate_groups"] if g["kind"] == "repeated_value"]
        assert vals and vals[0]["value"] == 413.762
        assert vals[0]["severity"] == "suspicious"
        assert vals[0]["n_columns"] == 2

    def test_cross_sheet_recycled_rows(self, tmp_path):
        rng = random.Random(11)
        wb = openpyxl.Workbook()
        ws1 = wb.active
        ws1.title = "Study1"
        ws1.append(["subject", "score_a", "score_b"])
        study1 = []
        for i in range(30):
            row = [f"s{i}", round(rng.uniform(1, 7), 3), round(rng.uniform(10, 99), 2)]
            study1.append(row)
            ws1.append(row)
        ws2 = wb.create_sheet("Study2")
        ws2.append(["subject", "score_a", "score_b"])
        for i in range(30):
            if 10 <= i < 14:
                ws2.append([f"t{i}", study1[i][1], study1[i][2]])
            else:
                ws2.append([f"t{i}", round(rng.uniform(1, 7), 3), round(rng.uniform(10, 99), 2)])
        p = tmp_path / "cross.xlsx"
        wb.save(p)
        res = detect_raw_duplication(str(p))
        assert any(g.get("cross_sheet") for g in res["duplicate_groups"])
        assert res["severity"] in ("suspicious", "highly_suspicious")

    def test_severity_never_exceeds_highly_suspicious(self, tmp_path):
        # A whole block of rows duplicated wholesale: strongest possible case.
        rows = _clean_rows()
        for k in range(10):
            src = rows[1 + k]
            rows[31 + k] = [31 + k, src[1], src[2], src[3]]
        res = detect_raw_duplication(str(_write_csv(tmp_path / "w.csv", rows)))
        assert res["severity"] == "highly_suspicious"  # capped, never impossible


# ---------------------------------------------------------------------------
# false-positive guards
# ---------------------------------------------------------------------------


class TestGuards:
    def test_clean_data_is_consistent(self, clean_csv):
        res = detect_raw_duplication(str(clean_csv))
        assert res["severity"] == "consistent"
        assert res["n_duplicate_groups"] == 0

    def test_id_column_excluded_as_arithmetic_run(self, clean_csv):
        res = detect_raw_duplication(str(clean_csv))
        assert any("arithmetic run" in c for c in res["excluded_columns"])

    def test_constant_column_excluded(self, tmp_path):
        rows = [["id", "dose", "mass"]]
        rng = random.Random(3)
        for i in range(1, 31):
            rows.append([i, 12.5, round(rng.uniform(1, 9), 3)])
        res = detect_raw_duplication(str(_write_csv(tmp_path / "c.csv", rows)))
        assert any("constant" in c for c in res["excluded_columns"])
        assert res["severity"] == "consistent"

    def test_low_cardinality_row_pairs_are_chance_not_findings(self, tmp_path):
        """Two shared columns with few distinct values produce many pairs
        by pigeonhole. The chance gate must CLEAR them (independence is
        only used to drop, never to escalate)."""
        rows = [["id", "arm", "visit", "score"]]
        rng = random.Random(11)
        for i in range(1, 41):
            rows.append([i, (i % 4) + 1, (i % 5) + 1, round(rng.uniform(1, 9), 3)])
        res = detect_raw_duplication(str(_write_csv(tmp_path / "arms.csv", rows)))
        assert not [g for g in res["duplicate_groups"] if g["kind"] == "row_pair"]

    def test_a_unique_high_precision_copy_still_fires(self, tmp_path):
        rows = [["id", "weight_kg", "cortisol_nmol"]]
        rng = random.Random(13)
        for i in range(1, 21):
            rows.append([i, round(rng.uniform(52, 96), 2), round(rng.uniform(180, 620), 3)])
        rows.append([21, rows[1][1], rows[1][2]])
        res = detect_raw_duplication(str(_write_csv(tmp_path / "copy.csv", rows)))
        pairs = [g for g in res["duplicate_groups"] if g["kind"] == "row_pair"]
        assert pairs, "a two-row copy of unique high-precision values must fire"

    def test_shared_covariate_never_accuses(self, tmp_path):
        # The canonical weather case: one site-level 7-digit value carried by
        # 20 rows. It must not produce row-pair findings, and the repeated
        # value itself must stay at indeterminate or below.
        rng = random.Random(5)
        rows = [["id", "site", "rainfall_mm", "mass_g"]]
        for i in range(1, 41):
            site = 1 if i <= 20 else 2
            rain = 123.4567 if site == 1 else 88.1234
            rows.append([i, site, rain, round(rng.uniform(1, 9), 3)])
        res = detect_raw_duplication(str(_write_csv(tmp_path / "wx.csv", rows)))
        assert not [g for g in res["duplicate_groups"] if g["kind"] == "row_pair"]
        order = ["highly_suspicious", "suspicious", "indeterminate", "consistent"]
        assert order.index(res["severity"]) >= order.index("indeterminate")
        singles = [g for g in res["duplicate_groups"] if g["kind"] == "repeated_value"]
        assert all(g["single_column_repeat"] for g in singles)

    def test_exclude_columns_silences_shared_covariate(self, tmp_path):
        rng = random.Random(5)
        rows = [["id", "site", "rainfall_mm", "mass_g"]]
        for i in range(1, 41):
            rows.append(
                [
                    i,
                    1 if i <= 20 else 2,
                    123.4567 if i <= 20 else 88.1234,
                    round(rng.uniform(1, 9), 3),
                ]
            )
        res = detect_raw_duplication(
            str(_write_csv(tmp_path / "wx.csv", rows)),
            exclude_columns=["rainfall_mm"],
        )
        assert res["severity"] == "consistent"
        assert any("shared-by-design" in c for c in res["excluded_columns"])

    def test_transform_column_precision_is_discounted(self, tmp_path):
        # sqrt of 2-dp values prints with ~10 decimals; a duplicated pair in
        # that column carries the RADICAND's information, not 10 digits'.
        rng = random.Random(9)
        rows = [["id", "raw", "sqrt_raw"]]
        for i in range(1, 31):
            raw_v = round(rng.uniform(1, 40), 2)
            rows.append([i, raw_v, round(math.sqrt(raw_v), 10)])
        rows[10][2] = rows[25][2]
        res = detect_raw_duplication(str(_write_csv(tmp_path / "t.csv", rows)))
        assert res["severity"] == "consistent"
        assert any("transformed" in c for c in res.get("caveats", []))

    def test_year_column_repeats_are_not_evidence(self, tmp_path):
        rng = random.Random(13)
        rows = [["id", "year", "mass"]]
        years = [1995, 2003, 2003, 2011, 1995, 2020, 2011, 2003] * 4
        for i, y in enumerate(years, 1):
            rows.append([i, y, round(rng.uniform(1, 9), 3)])
        res = detect_raw_duplication(str(_write_csv(tmp_path / "y.csv", rows)))
        assert res["severity"] == "consistent"

    def test_zeros_never_match(self, tmp_path):
        rng = random.Random(17)
        rows = [["id", "a", "b", "c"]]
        for i in range(1, 31):
            rows.append([i, 0.0, 0.0, round(rng.uniform(100, 999), 3)])
        res = detect_raw_duplication(str(_write_csv(tmp_path / "z.csv", rows)))
        assert not [g for g in res["duplicate_groups"] if g["kind"] == "row_pair"]

    def test_chance_level_repeats_not_reported(self, tmp_path):
        # 2000x4 grid of 2-dp values in [0, 1000): birthday collisions are
        # CERTAIN and must be absorbed by the Poisson gate.
        rng = random.Random(23)
        rows = [["a", "b", "c", "d"]]
        for _ in range(2000):
            rows.append([round(rng.uniform(0, 1000), 2) for _ in range(4)])
        res = detect_raw_duplication(str(_write_csv(tmp_path / "big.csv", rows)))
        order = ["highly_suspicious", "suspicious", "indeterminate", "consistent"]
        assert order.index(res["severity"]) >= order.index("indeterminate"), (
            "chance-level collisions in a large random grid were reported as "
            f"evidence: {res['severity']}"
        )

    def test_pooled_context_caps_at_indeterminate(self, tmp_path):
        rows = _clean_rows()
        rows[47] = [47, rows[12][1], rows[12][2], rows[12][3]]
        res = detect_raw_duplication(
            str(_write_csv(tmp_path / "p.csv", rows)),
            data_context="pooled_or_meta_analysis",
        )
        assert res["severity"] == "indeterminate"
        assert all(
            g["severity"] in ("indeterminate", "consistent") for g in res["duplicate_groups"]
        )
        assert any("by construction" in c.lower() or "BY CONSTRUCTION" in c for c in res["caveats"])


# ---------------------------------------------------------------------------
# structure and refusals
# ---------------------------------------------------------------------------


class TestStructure:
    def test_stacked_tables_are_segmented(self, tmp_path):
        rng = random.Random(29)
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["Table S1: control"])
        ws.append(["id", "mass"])
        for i in range(1, 11):
            ws.append([i, round(rng.uniform(5, 25), 3)])
        ws.append([])
        ws.append([])
        ws.append(["id", "length"])
        for i in range(1, 11):
            ws.append([i, round(rng.uniform(100, 300), 3)])
        p = tmp_path / "stacked.xlsx"
        wb.save(p)
        res = detect_raw_duplication(str(p))
        assert res["n_tables"] == 2
        assert any("stacks" in c for c in res["caveats"])

    def test_sheet_selector_and_missing_sheet(self, tmp_path):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "A"
        ws.append(["x"])
        ws.append([1.234])
        ws.append([5.678])
        wb.create_sheet("B").append(["y"])
        p = tmp_path / "two.xlsx"
        wb.save(p)
        assert detect_raw_duplication(str(p), sheet="A")["n_tables"] == 1
        err = detect_raw_duplication(str(p), sheet="Nope")
        assert err["severity"] == "indeterminate" and "not found" in err["error"]

    def test_missing_file_refuses(self):
        res = detect_raw_duplication("/nonexistent/supplement.xlsx")
        assert res["severity"] == "indeterminate"
        assert "not found" in res["error"]

    def test_unreadable_format_refuses_with_named_reason(self, tmp_path):
        p = tmp_path / "legacy.doc"
        p.write_bytes(b"\xd0\xcf\x11\xe0old word file")
        res = detect_raw_duplication(str(p))
        assert res["severity"] == "indeterminate"
        assert ".doc" in res["error"] or "doc" in res["error"]

    def test_invalid_data_context_refuses(self, clean_csv):
        res = detect_raw_duplication(str(clean_csv), data_context="banana")
        assert res["severity"] == "indeterminate"
        assert "data_context" in res["error"]

    def test_empty_file_refuses(self, tmp_path):
        p = _write_csv(tmp_path / "empty.csv", [["only", "headers"]])
        res = detect_raw_duplication(str(p))
        assert res["severity"] == "indeterminate"
        assert "no numeric data" in res["error"]

    def test_row_pair_budget_truncation_is_announced(self, tmp_path):
        rows = _clean_rows()
        rows[47] = [47, rows[12][1], rows[12][2], rows[12][3]]
        rows[50] = [50, rows[20][1], rows[20][2], rows[20][3]]
        res = detect_raw_duplication(
            str(_write_csv(tmp_path / "b.csv", rows)),
            max_row_pairs=1,
        )
        assert any("PARTIAL" in c for c in res["caveats"])

    def test_result_is_json_serialisable(self, tmp_path):
        import json

        rows = _clean_rows()
        rows[47] = [47, rows[12][1], rows[12][2], rows[12][3]]
        res = detect_raw_duplication(str(_write_csv(tmp_path / "j.csv", rows)))
        json.dumps(res)


class TestRecalibration:
    """Guards added after a 33-file corpus sweep flagged 21 of 24 files.

    Every case below is a real false positive that reached
    `highly_suspicious`, reduced to the smallest fixture that reproduces it.
    """

    def test_a_delimited_list_is_not_a_number(self):
        """THE PARSER BUG. Commas were stripped unconditionally, so the
        Qualtrics task order "1,2,4,3" became the integer 1243 -- four digits
        of precision invented out of a 4-item permutation. Ten columns of one
        corpus file were numeric only because of this, and they produced 48%
        of its findings."""
        from copy_paste_detector.readers import _as_number

        assert _as_number("1,2,4,3") is None
        assert _as_number("8,5,6,7") is None
        assert _as_number("1,23") is None
        # ...but a real thousands separator still reads as a number.
        assert _as_number("1,234") == 1234.0
        assert _as_number("1,234.56") == 1234.56
        assert _as_number("-1,000") == -1000.0

    def test_a_categorical_code_repeated_thousands_of_times_is_not_evidence(self, tmp_path):
        """THE POISSON INVERSION. The chance gate rejects a 4-digit value seen
        twice, but `expected_repeats` collapses to 0.0 in the far tail -- so
        the SAME value seen thousands of times passed, and scored HIGHER
        (score grows with log2(count)). Massive repetition is the strongest
        evidence the uniform-digit model does not fit the column."""
        rng = random.Random(5)
        rows = [["id", "condition_code", "score"]]
        codes = [1243, 1432, 1234, 5687, 5876, 1342]
        for i in range(1, 601):
            rows.append([i, codes[i % len(codes)], round(rng.uniform(1, 9), 3)])
        res = detect_raw_duplication(str(_write_csv(tmp_path / "cat.csv", rows)))
        order = ["highly_suspicious", "suspicious", "indeterminate", "consistent"]
        assert order.index(res["severity"]) >= order.index("indeterminate")

    def test_instrument_telemetry_corroborates_but_never_accuses(self, tmp_path):
        """Two participants on one Chrome build who both clicked once is the
        survey platform repeating itself, not a duplicated record -- so
        telemetry may not SEED a pair or stand alone as a repeated value.

        But it is not excluded either. Timers are the most identifying numbers
        in a Qualtrics export (a page-submit time of 226.971s is nobody
        else's), and excluding them outright cost real sensitivity: three
        participant rows copied into a corpus file became invisible, because
        22 of the high-information cells the duplicate shared were timers."""
        rng = random.Random(6)
        rows = [["id", "Q103_Timer_Page Submit", "Browser Version", "score"]]
        for i in range(1, 41):
            # Varying, so the constant-column screen does not claim them first.
            rows.append(
                [
                    i,
                    3.681 if i % 3 else round(rng.uniform(2, 9), 3),
                    18.18363 if i % 4 else 18.17763,
                    round(rng.uniform(1, 9), 3),
                ]
            )
        res = detect_raw_duplication(str(_write_csv(tmp_path / "t.csv", rows)))
        joined = " ".join(res["corroborating_columns"])
        assert "Page Submit" in joined and "Browser Version" in joined
        # Corroborating, NOT excluded: it may support a pair, never raise one.
        assert not any("Page Submit" in c for c in res["excluded_columns"])

    def test_a_numeric_minority_in_a_text_column_is_excluded(self, tmp_path):
        """Browser builds are strings ("80.0.3987.163"), but the few shaped
        like "75.0" parse as floats -- leaving a numeric minority that reads
        as high-precision data. The column is text; drop it whole."""
        rng = random.Random(8)
        rows = [["id", "build", "score"]]
        for i in range(1, 41):
            build = "75.0" if i % 10 == 0 else f"80.0.3987.{100 + i}"
            rows.append([i, build, round(rng.uniform(1, 9), 3)])
        res = detect_raw_duplication(str(_write_csv(tmp_path / "m.csv", rows)))
        assert any("numeric minority" in c for c in res["excluded_columns"])

    def test_a_repeated_value_finding_names_its_columns(self, tmp_path):
        """654 of 783 groups in the corpus sweep arrived with a cell
        reference and no header, so judging one meant opening the file and
        mapping a column index by hand."""
        rows = _clean_rows()
        rows[5][3] = 413.762
        rows[33][2] = 413.762
        res = detect_raw_duplication(str(_write_csv(tmp_path / "n.csv", rows)))
        vals = [g for g in res["duplicate_groups"] if g["kind"] == "repeated_value"]
        assert vals and vals[0]["columns"] == ["cortisol_nmol", "weight_kg"]
        assert any(e["column"] for e in res["by_column"])

    def test_model_output_is_capped_even_when_called_raw(self, tmp_path):
        """Six `tab_s*_model_results.xlsx` files were scanned as participant
        data. An R-squared repeats across a model's rows BY CONSTRUCTION."""
        rng = random.Random(9)
        rows = [["Term", "Beta", "SE", "Tvalue", "PValue", "R_Squared"]]
        for i in range(28):
            rows.append(
                [
                    f"term_{i}",
                    round(rng.uniform(-3, 3), 6),
                    round(rng.uniform(0.1, 2), 6),
                    round(rng.uniform(-4, 4), 6),
                    # the BH plateau: a clamped ceiling repeating
                    0.9865261522832445 if i % 7 >= 4 else round(rng.uniform(0, 1), 6),
                    0.7654321,
                ]
            )
        res = detect_raw_duplication(
            str(_write_csv(tmp_path / "mod.csv", rows)), data_context="raw_participant_level"
        )
        order = ["highly_suspicious", "suspicious", "indeterminate", "consistent"]
        assert order.index(res["severity"]) >= order.index("indeterminate")
        assert any("MODEL OUTPUT" in c for c in res["caveats"])

    def test_an_oversized_docx_is_read_in_full_not_as_a_preview(self, tmp_path):
        """`extract_supplement` caps its payload at 120k chars and returns an
        INDEX instead -- 5 rows x 12 columns per table, for a reader with a
        context window to choose from. This tool reads the grid in code, so it
        must go to the raw extractor: scanning a 5-row preview of a 17-row
        table and reporting the result as a verdict on the file is the
        truncated-sweep failure the coverage caps exist to prevent.

        A 263k-character corpus supplement hit exactly this and raised
        KeyError('data') -- loud, and lucky.
        """
        docx = pytest.importorskip("docx")
        path = tmp_path / "big.docx"
        doc = docx.Document()
        # Wide enough content to exceed the 120k index threshold.
        # Numbers in the first five columns; the bulk that pushes the file
        # past the index threshold goes in a text column, so the grid the
        # detector reads stays numeric.
        table = doc.add_table(rows=17, cols=6)
        for r in range(17):
            for c in range(5):
                table.cell(r, c).text = f"{r * 5 + c}.{r}{c}"
            table.cell(r, 5).text = "x" * 9000
        doc.save(str(path))

        res = detect_raw_duplication(str(path))
        assert "error" not in res, res.get("error")
        # 17 rows, not the 5 the index would have carried.
        assert res["n_numeric_cells"] > 5 * 6

    def test_planted_rows_survive_every_guard_on_a_real_survey_export(self, tmp_path):
        """THE SENSITIVITY CONTROL. Every other guard in this class REMOVES
        findings; without a positive control on the same shape of file, they
        could be suppressing the signal along with the noise, and the suite
        would look greener the more blind the tool became.

        Built to match a Qualtrics export: three metadata header rows, ~30
        low-information Likert columns that collide constantly between honest
        respondents, and a handful of high-precision page timers. Copy three
        complete participant records to new row positions and they must still
        be caught -- on the real corpus file this reproduces, the planted rows
        score 5.1-13.7 against a diffuse background that tops out at 1.37.

        This caught two real regressions: excluding telemetry outright made
        the duplicates invisible (the timers were the only columns precise
        enough to identify them), and the diffuse-evidence cap then demoted
        the survivors to indeterminate.
        """
        rng = random.Random(17)
        header = (
            ["pid"]
            + [f"q{i}" for i in range(30)]
            + ["Q1_Timer_Page Submit", "Q2_Timer_Page Submit", "Duration (in seconds)"]
        )
        body = []
        for i in range(1, 401):
            body.append(
                [i]
                # Likert 1-5: honest respondents collide here all the time.
                + [rng.randint(1, 5) for _ in range(30)]
                + [
                    round(rng.uniform(3, 300), 3),
                    round(rng.uniform(3, 300), 3),
                    rng.randint(100, 3000),
                ]
            )
        for src_i, dst_i in ((5, 150), (6, 250), (7, 350)):
            body[dst_i] = list(body[src_i])
            body[dst_i][0] = dst_i + 1  # a fresh participant id

        res = detect_raw_duplication(str(_write_csv(tmp_path / "q.csv", [header] + body)))
        pairs = [g for g in res["duplicate_groups"] if g["kind"] == "row_pair"]
        assert pairs, "three copied participant records were not detected"
        assert res["severity"] in ("suspicious", "highly_suspicious"), (
            f"planted duplicates were demoted to {res['severity']}"
        )

    def test_a_clean_survey_export_of_the_same_shape_stays_quiet(self, tmp_path):
        """The negative half of the control above: the same generator with
        nothing planted must not accuse. Otherwise the test above would pass
        on a tool that flags every Qualtrics file it sees."""
        rng = random.Random(17)
        header = (
            ["pid"]
            + [f"q{i}" for i in range(30)]
            + ["Q1_Timer_Page Submit", "Q2_Timer_Page Submit", "Duration (in seconds)"]
        )
        rows = [header]
        for i in range(1, 401):
            rows.append(
                [i]
                + [rng.randint(1, 5) for _ in range(30)]
                + [
                    round(rng.uniform(3, 300), 3),
                    round(rng.uniform(3, 300), 3),
                    rng.randint(100, 3000),
                ]
            )
        res = detect_raw_duplication(str(_write_csv(tmp_path / "c.csv", rows)))
        order = ["highly_suspicious", "suspicious", "indeterminate", "consistent"]
        assert order.index(res["severity"]) >= order.index("indeterminate")


# ---------------------------------------------------------------------------
# Identifier columns and repeated-measures shape
#
# Both classes below come from a sweep of the seven study CSVs behind PNAS
# 2217551120, where the detector's only two hits were artifacts of the file's
# SHAPE rather than of its contents. An accusation manufactured that way is the
# failure this module guards against everywhere else.
# ---------------------------------------------------------------------------


class TestGeographicCodesAreNotMeasurements:
    """A ZIP is a label written in numerals. Measured on `study1b_data.csv`:
    1,818 respondents across 565 distinct ZIPs produced 469 of the file's 471
    repeated-value groups, and drove the verdict to `suspicious`."""

    def test_zip_column_does_not_drive_a_finding(self, tmp_path):
        rng = random.Random(5)
        zips = [rng.choice([10009, 10022, 60602, 20175, 29681, 94110]) for _ in range(300)]
        rows = [["pid", "zip", "score"]]
        for i, z in enumerate(zips, start=1):
            rows.append([i, z, round(rng.uniform(1, 7), 3)])
        res = detect_raw_duplication(
            str(_write_csv(tmp_path / "z.csv", rows)), data_context="raw_participant_level"
        )
        assert not [g for g in res["duplicate_groups"] if "zip" in (g.get("columns") or [])]
        assert any("zip" in e for e in res["excluded_columns"])

    @pytest.mark.parametrize(
        "header", ["zip", "ZIP Code", "zip_code", "postcode", "FIPS", "county_code"]
    )
    def test_common_spellings_are_recognised(self, header):
        from copy_paste_detector.screening import _is_geocode

        assert _is_geocode(header)

    @pytest.mark.parametrize(
        "header", ["zipf_rank", "zippy", "fipsize", "score", "tractable_index"]
    )
    def test_a_substring_match_does_not_drop_a_real_column(self, header):
        """An exclusion that hides a finding is as bad as a false one."""
        from copy_paste_detector.screening import _is_geocode

        assert not _is_geocode(header)


class TestLongFormatIsShapeNotDuplication:
    """In long format each subject owns k rows, so subject-level attributes are
    written k times BY CONSTRUCTION. Measured on `study1a_data.csv`: 1,650 rows
    = 550 respondents x 3 conditions, which presented as a perfect 550-long
    repeat and earned `highly_suspicious`."""

    @staticmethod
    def _long_rows(n_subj=120, k=3, seed=11):
        rng = random.Random(seed)
        subj = [(round(rng.uniform(100, 900), 2), rng.randint(18, 80)) for _ in range(n_subj)]
        rows = [["duration", "age", "condition", "rating", "confidence"]]
        for block in range(k):
            for dur, age in subj:
                rows.append([dur, age, block + 1, rng.randint(1, 7), rng.randint(1, 7)])
        return rows

    def test_repeated_measures_file_is_not_accused(self, tmp_path):
        res = detect_raw_duplication(
            str(_write_csv(tmp_path / "long.csv", self._long_rows())),
            data_context="raw_participant_level",
        )
        assert res["severity"] == "consistent", res.get("explanation")

    def test_the_reshape_is_named_not_silently_tolerated(self, tmp_path):
        res = detect_raw_duplication(
            str(_write_csv(tmp_path / "long.csv", self._long_rows())),
            data_context="raw_participant_level",
        )
        assert any("long format" in e for e in res["excluded_columns"])

    def test_varying_columns_are_still_compared(self, tmp_path):
        """A reshape is no reason to stop looking.

        The planted run goes in a HIGH-PRECISION varying column, because that
        is what this detector treats as evidence: a run of small integers in a
        7-point rating column is correctly discounted by the diversity guard
        whether or not the file is long format, so planting there would test
        nothing. Verified by planting the same run in a plain wide file.
        """
        rng = random.Random(23)
        subj = [(round(rng.uniform(100, 900), 2), rng.randint(18, 80)) for _ in range(120)]
        rows = [["duration", "age", "condition", "rating", "latency_ms"]]
        for block in range(3):
            for dur, age in subj:
                rows.append(
                    [dur, age, block + 1, rng.randint(1, 7), round(rng.uniform(200, 9000), 3)]
                )
        col = rows[0].index("latency_ms")
        for offset in range(10):
            rows[5 + offset][col] = rows[200 + offset][col] = round(1234.567 + offset * 13.081, 3)
        res = detect_raw_duplication(
            str(_write_csv(tmp_path / "planted.csv", rows)), data_context="raw_participant_level"
        )
        assert res["severity"] != "consistent", (
            "long-format handling suppressed a planted duplicate"
        )
        assert any("long format" in e for e in res["excluded_columns"]), (
            "the file should still be recognised as long format"
        )

    def test_constant_columns_cannot_establish_a_period(self, tmp_path):
        """A constant column repeats across ANY partition. Counting it as
        support found `long format` in `study5_data.csv`, whose only repeating
        columns were Progress=100 and Finished=1 -- evidence that fits 2x396
        and 4x198 equally well."""
        rng = random.Random(3)
        rows = [["pid", "finished", "progress", "score"]]
        for i in range(1, 301):
            rows.append([i, 1, 100, round(rng.uniform(1, 9), 3)])
        res = detect_raw_duplication(
            str(_write_csv(tmp_path / "k.csv", rows)), data_context="raw_participant_level"
        )
        assert not any("long format" in e for e in res["excluded_columns"])

    def test_a_wholly_duplicated_file_is_still_caught(self, tmp_path):
        """Everything repeating is not long format -- it is a copied file."""
        rng = random.Random(9)
        body = [
            [i, round(rng.uniform(10, 99), 3), round(rng.uniform(10, 99), 3)] for i in range(1, 61)
        ]
        rows = [["pid", "a", "b"]] + body + body + body
        res = detect_raw_duplication(
            str(_write_csv(tmp_path / "dup.csv", rows)), data_context="raw_participant_level"
        )
        assert res["severity"] != "consistent"


class TestTelemetryNameLocatesAndTheFileJudges:
    """ "Duration" is a survey platform's page-timer AND a clinical measurement.
    The screen matched the word and demoted the column to corroborating-only,
    which bars it from ever SEEDING a duplicate pair -- so in a duration-heavy
    clinical dataset the highest-information columns were removed from the
    evidence base. The name now only locates a candidate; whether the file is a
    platform export decides it.
    """

    QUALTRICS = [
        "ResponseId",
        "StartDate",
        "EndDate",
        "Progress",
        "Duration (in seconds)",
        "Finished",
        "RecordedDate",
        "Q1",
        "Q2",
    ]
    CLINICAL = [
        "PatientID",
        "Sleep duration (h)",
        "Symptom Duration (days)",
        "disease_progression_score",
        "MMSE version B",
        "Exercise duration min",
        "Age",
        "Sex",
    ]

    def test_clinical_durations_are_not_platform_telemetry(self):
        from copy_paste_detector.screening import _is_telemetry

        fooled = [
            n
            for n in (
                "Sleep duration (h)",
                "Symptom Duration (days)",
                "Exercise duration min",
                "MMSE version B",
            )
            if _is_telemetry(n, self.CLINICAL)
        ]
        assert not fooled, f"clinical measurements screened as telemetry: {fooled}"

    def test_the_same_names_ARE_telemetry_inside_a_platform_export(self):
        from copy_paste_detector.screening import _is_telemetry

        missed = [
            n
            for n in ("Duration (in seconds)", "Progress", "Finished", "StartDate")
            if not _is_telemetry(n, self.QUALTRICS)
        ]
        assert not missed, f"real telemetry no longer screened: {missed}"

    def test_platform_only_names_need_no_context(self):
        from copy_paste_detector.screening import _is_telemetry

        for n in ("First Click", "Distribution Channel", "User Agent", "IP Address", "Page Submit"):
            assert _is_telemetry(n, self.CLINICAL), n

    def test_without_context_only_the_unambiguous_names_count(self):
        """The safe reading: demoting a column removes its values from the
        evidence base, so an ambiguous name with nothing to judge it by is
        kept."""
        from copy_paste_detector.screening import _is_telemetry

        assert _is_telemetry("First Click") is True
        assert _is_telemetry("Duration (in seconds)") is False

    def test_word_bounded_so_progression_is_not_progress(self):
        from copy_paste_detector.screening import _telemetry_name_candidate

        assert _telemetry_name_candidate("disease_progression_score") is None
        assert _telemetry_name_candidate("Progress") == "ambiguous"

    def test_camelCase_headers_normalise(self):
        """Not cosmetic: without this `ResponseId` fails the platform test, a
        real Qualtrics export is not recognised as one, and its ambiguous
        columns are judged with no context at all."""
        from copy_paste_detector.screening import _looks_like_platform_export, _normalise_header

        assert _normalise_header("ResponseId") == "response id"
        assert _looks_like_platform_export(self.QUALTRICS) is True
        assert _looks_like_platform_export(self.CLINICAL) is False

    def test_the_duplicated_rows_are_FOUND_not_merely_uncapped(self, tmp_path):
        """The consequence, not the mechanism. Two copied participant rows in a
        clinical file with no platform columns: the old screen demoted all
        three duration columns, nothing was left to seed a pair, and the tool
        reported "No duplicated raw data detected"."""
        import csv
        import random

        from copy_paste_detector import detect_raw_duplication

        random.seed(11)
        hdr = [
            "PatientID",
            "Sleep duration (h)",
            "Symptom Duration (days)",
            "Exercise duration min",
            "Age",
        ]
        rows = [
            [
                f"P{i:03d}",
                round(random.uniform(4, 9), 2),
                round(random.uniform(1, 300), 1),
                round(random.uniform(0, 120), 1),
                random.randint(40, 85),
            ]
            for i in range(60)
        ]
        rows.append(["P900"] + rows[7][1:])
        rows.append(["P901"] + rows[23][1:])
        p = tmp_path / "clinical_dup.csv"
        with open(p, "w", newline="") as fh:
            w = csv.writer(fh)
            w.writerow(hdr)
            w.writerows(rows)

        r = detect_raw_duplication(str(p))
        assert r["n_duplicate_groups"] == 2, r["explanation"]
        assert r["corroborating_columns"] == []
        # And the overruled screen is ANNOUNCED, not silently dropped.
        assert any("telemetry screen by NAME" in c for c in r["caveats"])


class TestTheCountsAreNamedForWhatTheyCount:
    """ "75 duplicated row pair(s)" was a count of GROUPS, and got divided by rows.

    MEASURED on a 193-row, 2-sheet workbook: the tool reported "75 duplicated
    row pair(s)", a reader divided 75 by 193 to get "roughly 39% of rows
    implicated", and both numbers reached a draft accusation. 75 was the number
    of merged match-signature GROUPS across BOTH sheets; the strict full-row
    figure for one sheet was 22 pairs over 24 rows, i.e. 12.4%. Three different
    objects had been sharing one noun, and nothing in the payload let a caller
    tell them apart.
    """

    def test_the_three_counts_are_separate_keys(self, tmp_path):
        rows = _clean_rows(40)
        rows.append(list(rows[5]))  # one planted duplicate row
        res = detect_raw_duplication(str(_write_csv(tmp_path / "d.csv", rows)))
        for k in (
            "n_row_pair_groups",
            "n_row_pairs",
            "n_rows_implicated",
            "n_rows_total",
            "pct_rows_implicated",
        ):
            assert k in res, k
        # A group merges pairs, so it can never exceed the pair count.
        assert res["n_row_pair_groups"] <= res["n_row_pairs"]
        assert res["n_rows_implicated"] <= res["n_rows_total"]

    def test_the_percentage_is_over_rows_not_groups(self, tmp_path):
        rows = _clean_rows(40)
        rows.append(list(rows[5]))
        res = detect_raw_duplication(str(_write_csv(tmp_path / "d.csv", rows)))
        expected = round(100.0 * res["n_rows_implicated"] / res["n_rows_total"], 1)
        assert res["pct_rows_implicated"] == expected
        assert res["pct_rows_implicated"] <= 100.0

    def test_the_explanation_says_group_not_pair(self, tmp_path):
        rows = _clean_rows(40)
        rows.append(list(rows[5]))
        res = detect_raw_duplication(str(_write_csv(tmp_path / "d.csv", rows)))
        assert "row-pair group" in res["explanation"]
        assert "not of pairs" in res["explanation"]

    def test_truncation_says_which_way_it_is_wrong(self, tmp_path):
        # Same shape as test_row_pair_budget_truncation_is_announced, which is
        # the setup known to reach the truncation branch.
        rows = _clean_rows()
        rows[47] = [47, rows[12][1], rows[12][2], rows[12][3]]
        rows[50] = [50, rows[20][1], rows[20][2], rows[20][3]]
        res = detect_raw_duplication(str(_write_csv(tmp_path / "d.csv", rows)), max_row_pairs=1)
        joined = " ".join(res["caveats"])
        assert "PARTIAL" in joined
        assert "UNDER-count" in joined, "a floor must be named as a floor"


class TestAColumnComputedFromTwoOthersIsNotIndependentEvidence:
    """One coincidence must not be counted three times.

    MEASURED on the CARS workbook: `QALY - SG` is exactly
    0.125*(`SG - V2`) + 0.125*(`SG - V7`) across all 193 rows, and the ORD and
    Bpost triples have the same shape. A chance tie on the two inputs FORCES a
    17-significant-digit tie in the output, which the scorer read as a third
    independent high-information match. The transform detector could not see
    it: it tests each column against a library of fractions/sqrt/log and has no
    notion of a column computed FROM two others.
    """

    @staticmethod
    def _rows(n=60, seed=11):
        rng = random.Random(seed)
        out = [["id", "pre", "post", "mean_qaly"]]
        for i in range(1, n + 1):
            # low-cardinality inputs, exactly as a utility mapping produces
            pre = round(rng.choice([0.61, 0.709, 0.652, 0.77, 0.557]), 3)
            post = round(rng.choice([0.852, 0.795, 0.658, 0.616, 0.618]), 3)
            out.append([i, pre, post, (pre + post) / 2 * 0.25])
        return out

    def test_the_derived_column_is_named_and_discounted(self, tmp_path):
        res = detect_raw_duplication(str(_write_csv(tmp_path / "d.csv", self._rows())))
        notes = " ".join(res["caveats"])
        assert "deterministic function" in notes
        assert "mean_qaly" in notes

    def test_a_two_level_column_is_not_called_derived(self, tmp_path):
        """A 3-parameter fit through <= 3 distinct points is exact by
        construction. `Treatment Group` (1/2) came out as a "function of"
        `Treatment cost` (two values) -- true, degenerate, not a formula."""
        rng = random.Random(3)
        rows = [["id", "arm", "cost", "score"]]
        for i in range(1, 61):
            arm = (i % 2) + 1
            rows.append([i, arm, 2104.2 if arm == 1 else 207.9, round(rng.uniform(1, 90), 3)])
        res = detect_raw_duplication(str(_write_csv(tmp_path / "b.csv", rows)))
        notes = " ".join(res["caveats"])
        assert "'arm' in" not in notes or "deterministic function" not in notes

    def test_honest_correlated_columns_are_not_called_derived(self, tmp_path):
        """Tight correlation is not a formula: only exact residuals count."""
        rng = random.Random(5)
        rows = [["id", "height", "weight", "bmi_ish"]]
        for i in range(1, 61):
            h = round(rng.uniform(1.5, 2.0), 3)
            w = round(rng.uniform(50, 100), 2)
            rows.append([i, h, w, round(w / (h * h) + rng.uniform(-2, 2), 3)])
        res = detect_raw_duplication(str(_write_csv(tmp_path / "c.csv", rows)))
        assert "bmi_ish" not in " ".join(res["caveats"])


def test_headerless_word_table_keeps_first_measurement_row(tmp_path):
    from docx import Document

    doc = Document()
    table = doc.add_table(rows=3, cols=3)
    for r, row in enumerate(table.rows):
        for c, cell in enumerate(row.cells):
            cell.text = str(12.345 + r * 7 + c * 0.12)
    path = tmp_path / "continuation.docx"
    doc.save(path)
    result = detect_raw_duplication(str(path))
    assert "error" not in result
    assert result["n_numeric_cells"] == 9
    assert result["n_rows_total"] == 3

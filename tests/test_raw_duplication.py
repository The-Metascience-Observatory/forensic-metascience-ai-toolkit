"""Agent compatibility and Word-reader parity for the standalone detector.

Algorithm and false-positive regressions live in copy_paste_detector/tests.
"""

from __future__ import annotations

import csv
import inspect
import json
import random

import pytest

from copy_paste_detector import detect_raw_duplication
from copy_paste_detector.docx_reader import extract_docx_tables


def test_the_tool_spec_signature_is_stable():
    assert list(inspect.signature(detect_raw_duplication).parameters) == [
        "path",
        "sheet",
        "exclude_columns",
        "data_context",
        "max_row_pairs",
        "limits",
    ]
    assert inspect.signature(detect_raw_duplication).parameters["limits"].kind is inspect.Parameter.KEYWORD_ONLY
    inspect.signature(detect_raw_duplication).bind("data.csv", None, [], "unknown", 5000)


@pytest.mark.parametrize(
    "shape", ["ordinary", "merged", "stacked", "blank", "header_only", "empty", "no_header"]
)
def test_word_input_grid_matches_agent_reader(tmp_path, shape):
    from docx import Document

    from metascience_agent.extraction.supplements_read import extract_docx_tables as agent_reader

    doc = Document()
    table = doc.add_table(rows=6 if shape != "header_only" else 1, cols=4)
    for r, row in enumerate(table.rows):
        for c, cell in enumerate(row.cells):
            cell.text = f"measurement {c}" if r == 0 else f"{r * 9 + c}.763"
    if shape == "no_header":
        for c, cell in enumerate(table.rows[0].cells):
            cell.text = f"{20 + c}.735"
    if shape == "merged":
        table.cell(2, 1).merge(table.cell(2, 2)).text = "12.735"
        table.cell(3, 0).merge(table.cell(4, 0)).text = "subject A"
    elif shape == "stacked":
        table.cell(0, 0).merge(table.cell(0, 3)).text = "Table 1. Measurements"
        table.cell(1, 0).merge(table.cell(1, 1)).text = "Group A"
        table.cell(1, 2).merge(table.cell(1, 3)).text = "Group B"
        for c in range(4):
            table.cell(2, c).text = f"Value {c}"
    elif shape in {"blank", "empty"}:
        for r in range(6) if shape == "empty" else [0, 3]:
            for cell in table.rows[r].cells:
                cell.text = " \n "
    path = tmp_path / "data.docx"
    doc.save(path)
    expected, actual = agent_reader(path), extract_docx_tables(path)
    assert actual["n_tables"] == expected["n_tables"]
    for old, new in zip(expected["tables"], actual["tables"]):
        assert new["label"] == old["label"]
        assert new["rows"] == old["rows"]
        assert new["data"] == old["data"]


def test_existing_agent_cli_and_registry_use_same_result(tmp_path):
    from typer.testing import CliRunner

    from metascience_agent.cli import app
    from metascience_agent.registry import BY_NAME

    rng = random.Random(7)
    rows = [["id", "weight", "cortisol"]]
    rows += [[i, round(rng.uniform(52, 96), 2), round(rng.uniform(180, 620), 3)] for i in range(60)]
    rows[47][1:] = rows[12][1:]
    path, output = tmp_path / "data.csv", tmp_path / "result.json"
    with path.open("w", newline="") as handle:
        csv.writer(handle).writerows(rows)
    expected = detect_raw_duplication(str(path))
    spec = BY_NAME["detect_raw_duplication"]
    assert spec.fn is detect_raw_duplication
    assert spec.fn(str(path)) == expected
    result = CliRunner().invoke(app, ["raw-duplication-detect", str(path), "--output", str(output)])
    assert result.exit_code == 0, result.output
    assert json.loads(output.read_text()) == expected

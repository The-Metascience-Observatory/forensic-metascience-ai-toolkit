"""A supplement too big to hand back whole is INDEXED, not refused.

THE FAILURE THIS CLOSES. On 10.1038/s41537-025-00576-3 the supplement held
eTable 5 -- the latent-class model-selection statistics (BIC and the other fit
indices) that justify the four-class solution the ENTIRE paper rests on.
Extraction succeeded, the result was 1,114,127 characters, it spilled to a temp
file, and the review recorded: "reading it back by line offset would sample
rather than extract, so no check was run."

MEASURED across every spreadsheet supplement in the working corpus:
    fnins.2023.1186751   2 sheets, 1,193 rows  ->  1,879,198 chars
    pone.0254493         3 sheets,   282 rows  ->  1,383,602 chars
    healthcare11101497   2 sheets,   386 rows  ->    498,993 chars
All three exceed any plausible single-result limit, so this was never one
unlucky paper: EVERY spreadsheet supplement was unreadable for the same reason.

The paper path solved this long ago -- `extract_xml` returns an index and
`get_table` fetches one table. Supplements had no equivalent.
"""
import json

import openpyxl
import pytest

from forensic_metascience_ai_toolkit.extraction.supplements_read import (
    _OVERSIZED_CHARS, extract_supplement, get_supplement_table)


@pytest.fixture
def big_book(tmp_path):
    """A workbook comfortably over the limit, with two sheets."""
    book = openpyxl.Workbook()
    s1 = book.active
    s1.title = "eTable5_fit"
    s1.append(["model", "classes", "BIC", "entropy"])
    for k in range(2, 8):
        s1.append([f"m{k}", k, 10000 - k * 37.5, 0.7 + k / 100])
    s2 = book.create_sheet("participants")
    s2.append([f"var{i}" for i in range(40)])
    for r in range(400):
        s2.append([f"r{r}c{i}" * 3 for i in range(40)])
    p = tmp_path / "supp.xlsx"
    book.save(p)
    return p


class TestTheIndexReplacesAnUnsendablePayload:
    def test_a_big_supplement_comes_back_as_an_index(self, big_book):
        r = extract_supplement(big_book)
        assert r["oversized"] is True
        assert len(json.dumps(r, default=str)) <= _OVERSIZED_CHARS
        assert r["full_size_chars"] > _OVERSIZED_CHARS

    def test_nothing_is_silently_dropped(self, big_book):
        """A truncated payload that LOOKED complete would be far worse than the
        refusal it replaces. Every sheet is still named, with its shape."""
        r = extract_supplement(big_book)
        assert r["n_tables"] == 2
        assert [t["label"] for t in r["tables"]] == ["eTable5_fit", "participants"]
        assert all(t["n_rows"] and t["width"] for t in r["tables"])

    def test_the_index_says_how_to_get_the_rest(self, big_book):
        r = extract_supplement(big_book)
        assert "get_supplement_table" in r["note"]
        assert "Do NOT treat this index as the supplement" in r["note"]

    def test_the_preview_is_enough_to_recognise_a_sheet_by(self, big_book):
        """A sheet called `Sheet1` is identified by what is IN it."""
        r = extract_supplement(big_book)
        fit = [t for t in r["tables"] if t["label"] == "eTable5_fit"][0]
        # The COLUMN NAMES are what distinguish a fit-indices sheet from a
        # participant listing, and the first version of the index left them out
        # because `rows` is the body. This test is why `headers` is there.
        assert "BIC" in fit["headers"]
        assert any("9925" in str(c) for row in fit["first_rows"] for c in row)

    def test_a_small_supplement_is_returned_whole(self, tmp_path):
        book = openpyxl.Workbook()
        book.active.append(["a", "b"])
        book.active.append(["1", "2"])
        p = tmp_path / "small.xlsx"
        book.save(p)
        r = extract_supplement(p)
        assert not r.get("oversized")
        assert r["tables"][0]["rows"] == [["1", "2"]]


class TestTheFetchWindowsEveryViewTogether:
    """THE BUG CAUGHT IN TESTING, and the reason this class exists. A table
    dict carries the same rows three times -- `rows`, `data` and `cells` -- and
    the first version sliced only `rows`. Asking for five rows of an 86-row
    sheet returned 1,252,595 characters, because `cells` still held all of
    them: the exact defect the window exists to prevent, one field over."""

    def test_a_window_actually_shrinks_the_result(self, big_book):
        whole = get_supplement_table(big_book, "participants")
        part = get_supplement_table(big_book, "participants", first_row=1, last_row=5)
        assert len(json.dumps(part, default=str)) < len(json.dumps(whole, default=str)) / 10

    def test_cells_are_windowed_with_the_rows(self, big_book):
        part = get_supplement_table(big_book, "participants", first_row=1, last_row=5)
        assert len(part["rows"]) == 5
        # `cells[].row` is 1-based into `data` = [header] + body, so body rows
        # 1..5 are cell rows 2..6. One shared coordinate is what makes this work.
        assert {c["row"] for c in part["cells"]} <= {2, 3, 4, 5, 6}

    def test_the_header_survives_every_window(self, big_book):
        """A window without its header is a grid of unlabelled numbers."""
        part = get_supplement_table(big_book, "participants", first_row=20, last_row=22)
        assert part["data"][0] == part["headers"]
        assert len(part["data"]) == 4

    def test_a_window_says_it_is_a_window(self, big_book):
        """A window that did not announce itself would be a silent truncation
        wearing the shape of a complete answer."""
        part = get_supplement_table(big_book, "participants", first_row=2, last_row=3)
        assert part["row_window"] == {"first_row": 2, "last_row": 3, "n_rows_total": 400}
        assert "say nothing either way" in part["window_note"]

    def test_an_unwindowed_fetch_carries_no_window_note(self, big_book):
        assert "window_note" not in get_supplement_table(big_book, "eTable5_fit")


class TestChoosingASheet:
    def test_by_number(self, big_book):
        assert get_supplement_table(big_book, 1)["label"] == "eTable5_fit"

    def test_by_name(self, big_book):
        """A caller that has just read an index holds both, and should not have
        to guess which one is wanted."""
        assert get_supplement_table(big_book, "participants")["table_num"] == 2

    def test_an_unknown_sheet_lists_the_real_ones(self, big_book):
        r = get_supplement_table(big_book, "eTable9")
        assert "no sheet" in r["error"]
        assert [t["label"] for t in r["available"]] == ["eTable5_fit", "participants"]


class TestItIsReachableFromTheModel:
    """Built-and-never-consumed is this repo's signature defect; a reader that
    nothing can call is the same bug as a reader that cannot read."""

    def test_the_tool_is_registered(self):
        from forensic_metascience_ai_toolkit import registry
        assert "get_supplement_table" in registry.BY_NAME

    def test_it_is_offered_to_both_stages(self):
        from forensic_metascience_ai_toolkit import registry
        spec = registry.BY_NAME["get_supplement_table"]
        assert spec.stages

"""Standalone execution, output artifacts and operational failure behavior."""

from __future__ import annotations

import csv
import json
import random
import subprocess
import sys
from pathlib import Path

import pytest

from copy_paste_detector import detect_raw_duplication
from copy_paste_detector.cli import main


@pytest.fixture(params=["csv", "tsv", "xlsx", "xlsm", "docx"])
def data_file(tmp_path, request):
    rng = random.Random(7)
    rows = [["id", "group", "weight_kg", "cortisol_nmol"]]
    rows += [
        [i, i % 3 + 1, round(rng.uniform(52, 96), 2), round(rng.uniform(180, 620), 3)]
        for i in range(1, 61)
    ]
    rows[47][1:] = rows[12][1:]
    path = tmp_path / f"data.{request.param}"
    if request.param in {"csv", "tsv"}:
        with path.open("w", newline="") as handle:
            csv.writer(handle, delimiter="\t" if request.param == "tsv" else ",").writerows(rows)
    elif request.param in {"xlsx", "xlsm"}:
        from openpyxl import Workbook

        book = Workbook()
        book.active.title = "Data"
        for row in rows:
            book.active.append(row)
        book.save(path)
        book.close()
    else:
        from docx import Document

        doc = Document()
        table = doc.add_table(rows=len(rows), cols=4)
        for row, values in zip(table.rows, rows):
            for cell, value in zip(row.cells, values):
                cell.text = str(value)
        doc.save(path)
    return path


def test_cli_matches_api_for_every_format(data_file, capsys):
    assert main([str(data_file)]) == 0
    captured = capsys.readouterr()
    assert not captured.err
    assert json.loads(captured.out) == detect_raw_duplication(str(data_file))


def test_module_runs_without_agent_imports(data_file, tmp_path):
    # -I discards inherited PYTHONPATH and user site packages. Add only this
    # project's package directory and refuse any attempted agent import.
    import copy_paste_detector

    source = str(Path(copy_paste_detector.__file__).resolve().parent.parent)
    script = """
import importlib.abc, runpy, sys
class NoAgent(importlib.abc.MetaPathFinder):
    def find_spec(self, fullname, path=None, target=None):
        if fullname == "forensic_metascience_ai_toolkit" or fullname.startswith("forensic_metascience_ai_toolkit."):
            raise AssertionError("Standalone execution attempted an agent import")
sys.meta_path.insert(0, NoAgent())
sys.path.insert(0, sys.argv.pop(1))
sys.argv[0] = "copy-paste-detect"
runpy.run_module("copy_paste_detector", run_name="__main__")
"""
    proc = subprocess.run(
        [sys.executable, "-I", "-c", script, source, str(data_file)],
        cwd=tmp_path,
        capture_output=True,
        text=True,
    )
    assert proc.returncode == 0, proc.stderr
    # The legacy detector chooses row representatives from sets, so both
    # evidence order and representative column can vary with the hash seed.
    # Same-process tests compare full dictionaries; here compare the findings.
    actual = json.loads(proc.stdout)
    expected = detect_raw_duplication(str(data_file))
    for key in [
        "severity",
        "n_row_pairs",
        "n_rows_implicated",
        "n_numeric_cells",
        "n_duplicate_groups",
    ]:
        assert actual[key] == expected[key]

    def matched_cells(result):
        return {
            (cell["sheet"], cell["cell"])
            for group in result["duplicate_groups"]
            for cell in group["cells"]
        }

    assert matched_cells(actual) == matched_cells(expected)


def test_json_and_highlight_outputs(data_file, tmp_path, capsys):
    output, highlight = tmp_path / "result.json", tmp_path / "highlight.xlsx"
    status = main([str(data_file), "-o", str(output), "--highlight-out", str(highlight)])
    result = json.loads(output.read_text())
    expected = detect_raw_duplication(str(data_file))
    assert result.pop("highlight") == (
        {
            "error": "cannot highlight .docx (no cell coordinates); only .xlsx/.xlsm/.csv/.tsv sources"
        }
        if data_file.suffix == ".docx"
        else {"path": str(highlight), "n_groups_highlighted": 1, "n_groups_skipped": 0}
    )
    assert result == expected
    assert status == (1 if data_file.suffix == ".docx" else 0)
    assert not capsys.readouterr().out


@pytest.mark.parametrize("kind", ["missing", "corrupt", "unsupported"])
def test_read_failures_are_nonzero_json_errors(tmp_path, capsys, kind):
    path = tmp_path / ("data.xlsx" if kind == "corrupt" else "data.xls")
    if kind != "missing":
        path.write_text("not a workbook")
    assert main([str(path)]) == 1
    captured = capsys.readouterr()
    result = json.loads(captured.out)
    assert result["severity"] == "indeterminate"
    assert result["error"]
    assert captured.err


def test_input_cannot_be_overwritten(data_file):
    before = data_file.read_bytes()
    for flag in ["--output", "--highlight-out"]:
        with pytest.raises(SystemExit) as exc:
            main([str(data_file), flag, str(data_file)])
        assert exc.value.code == 2
    assert data_file.read_bytes() == before


def test_output_failure_is_nonzero(data_file, tmp_path, capsys):
    assert main([str(data_file), "-o", str(tmp_path / "absent" / "result.json")]) == 1
    assert "Could not write JSON" in capsys.readouterr().err


def test_portable_run_reads_the_loaded_package_version(tmp_path):
    from copy_paste_detector import __version__
    from copy_paste_detector.portable import save_run
    output = save_run(tmp_path, 'cp_test', '2026-09-10', None, {'error': 'unavailable'}, {})
    assert json.loads((output / 'manifest.json').read_text())['version'] == __version__
